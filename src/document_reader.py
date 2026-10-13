import os
from pathlib import Path

from docx import Document
from PIL import Image, ImageEnhance, ImageOps, ImageSequence
import pytesseract

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}

def resolve_tesseract_executable() -> str | None:
    env_path = os.getenv("TESSERACT_CMD") or os.getenv("TESSERACT_PATH")
    if env_path and Path(env_path).exists():
        return env_path
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        str(Path.home() / "AppData" / "Local" / "Programs" / "Tesseract-OCR" / "tesseract.exe"),
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return None

def _configure_tesseract() -> None:
    resolved = resolve_tesseract_executable()
    if resolved:
        pytesseract.pytesseract.tesseract_cmd = resolved

def read_docx(path: str | Path) -> str:
    """Extract paragraphs and table cells from a DOCX document."""
    doc = Document(str(path))
    chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table_idx, table in enumerate(doc.tables, start=1):
        chunks.append(f"[TABLE {table_idx}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    # Include headers/footers when they contain useful agreement details.
    for section_idx, section in enumerate(doc.sections, start=1):
        for label, container in (("HEADER", section.header), ("FOOTER", section.footer)):
            text = "\n".join(p.text.strip() for p in container.paragraphs if p.text.strip())
            if text:
                chunks.append(f"[{label} {section_idx}]\n{text}")
    return "\n".join(chunks).strip()

def _preprocess_image(image: Image.Image) -> Image.Image:
    image = image.convert("L")
    image = ImageOps.autocontrast(image)
    return ImageEnhance.Contrast(image).enhance(1.4)

def read_image(path: str | Path) -> str:
    """OCR an image; Tesseract must be installed separately."""
    _configure_tesseract()
    chunks = []
    with Image.open(path) as image:
        # OCR all frames for multi-page TIFFs, not only the first frame.
        for frame in ImageSequence.Iterator(image):
            text = pytesseract.image_to_string(_preprocess_image(frame.copy()), config="--psm 3").strip()
            if text:
                chunks.append(text)
    return "\n\n".join(chunks).strip()

def read_document(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx(path)
    if suffix in IMAGE_SUFFIXES:
        return read_image(path)
    raise ValueError(
        f"Unsupported file type: {suffix}. Supported: DOCX and common image formats."
    )
